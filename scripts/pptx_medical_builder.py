#!/usr/bin/env python3
"""
pptx_medical_builder.py — Production Helper for Medical Journal PowerPoint (.pptx) Decks.

Addresses all architectural requirements:
- Native slide.background.fill (never a floating draggable full-slide rectangle).
- Populates hidden/accessible slide.shapes.title so Outline View and screen readers work.
- Strict Horizontal Bands (Title, Body, Footer) with collision buffers.
- Complete helpers: set_bullet, add_fragment_bullet, add_box, add_2_column, add_3_column_pillars,
  add_clinical_table, add_native_bar_chart, add_speaker_notes, crop_pdf_figure_300dpi, add_cropped_figure.
- Colorblind-safe palettes (Okabe-Ito inspired) + mandatory non-color symbols ([↑], [↓], [⚠], [✓]).
- Bilingual footer support (lang='id' -> 'Hal. X', lang='en' -> 'Slide X').
"""

import math
import os
from typing import List, Tuple, Optional, Dict
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from PIL import Image
import pypdfium2 as pdfium

# ── STRICT HORIZONTAL BANDS (13.333" x 7.5" Widescreen 16:9) ──
SLIDE_W    = 13.333
SLIDE_H    = 7.50
MARGIN_X   = 0.80
CONTENT_W  = 11.733
TITLE_TOP  = 0.45
TITLE_H    = 1.50   # Bottom = 1.95" (0.25" buffer before BODY_TOP)
BODY_TOP   = 2.20
BODY_H     = 4.55   # Bottom = 6.75" (0.20" buffer before FOOTER_TOP)
FOOTER_TOP = 6.95
FOOTER_H   = 0.35   # Bottom = 7.30"
COL_GAP    = 0.45
COL_W_3    = 3.61
COL_W_2    = 5.64

# ── CAPACITY BUDGET MATH ──
CHAR_W_EM = 0.52
LINE_H    = 1.22
RAGGED    = 0.90

# ── UNIFIED PALETTE PRESETS (Consistent across SKILL.md, AGENTS.md, README.md) ──
PALETTES: Dict[str, Dict[str, RGBColor]] = {
    "Clinical Navy": {
        "BG_THEME":   RGBColor(0x0F, 0x17, 0x2A),
        "BG_CONTENT": RGBColor(0xFF, 0xFF, 0xFF),
        "TEXT_LIGHT": RGBColor(0xF8, 0xFA, 0xFC),
        "TEXT_DARK":  RGBColor(0x0F, 0x17, 0x2A),
        "TEXT_MUTED": RGBColor(0x47, 0x55, 0x69),
        "ACCENT":     RGBColor(0x00, 0x72, 0xB2),  # Okabe-Ito Blue (Intervention)
        "CONTROL":    RGBColor(0x64, 0x74, 0x8B),  # Slate (Control/Placebo)
        "TINT":       RGBColor(0xEF, 0xF6, 0xFF),  # Soft Card Fill
        "ALERT_RED":  RGBColor(0xD5, 0x5E, 0x00),  # Okabe-Ito Vermilion (Colorblind-safe Alert)
        "POS_TEAL":   RGBColor(0x00, 0x9E, 0x73),  # Okabe-Ito Bluish Green (Colorblind-safe Positive)
        "TABLE_ALT":  RGBColor(0xF8, 0xFA, 0xFC),
    },
    "Deep Slate": {
        "BG_THEME":   RGBColor(0x1E, 0x29, 0x3B),
        "BG_CONTENT": RGBColor(0xFF, 0xFF, 0xFF),
        "TEXT_LIGHT": RGBColor(0xF8, 0xFA, 0xFC),
        "TEXT_DARK":  RGBColor(0x1E, 0x29, 0x3B),
        "TEXT_MUTED": RGBColor(0x47, 0x55, 0x69),
        "ACCENT":     RGBColor(0x25, 0x63, 0xEB),
        "CONTROL":    RGBColor(0x64, 0x74, 0x8B),
        "TINT":       RGBColor(0xF1, 0xF5, 0xF9),
        "ALERT_RED":  RGBColor(0xC5, 0x30, 0x30),
        "POS_TEAL":   RGBColor(0x0D, 0x94, 0x88),
        "TABLE_ALT":  RGBColor(0xF8, 0xFA, 0xFC),
    },
    "Medical Teal": {
        "BG_THEME":   RGBColor(0x11, 0x5E, 0x59),
        "BG_CONTENT": RGBColor(0xFF, 0xFF, 0xFF),
        "TEXT_LIGHT": RGBColor(0xF0, 0xFD, 0xFA),
        "TEXT_DARK":  RGBColor(0x13, 0x4E, 0x4A),
        "TEXT_MUTED": RGBColor(0x47, 0x55, 0x69),
        "ACCENT":     RGBColor(0x0D, 0x94, 0x88),
        "CONTROL":    RGBColor(0x64, 0x74, 0x8B),
        "TINT":       RGBColor(0xF0, 0xFD, 0xFA),
        "ALERT_RED":  RGBColor(0xD5, 0x5E, 0x00),
        "POS_TEAL":   RGBColor(0x0F, 0x76, 0x6E),
        "TABLE_ALT":  RGBColor(0xF0, 0xFD, 0xFA),
    },
    "Minimalist High-Contrast": {
        "BG_THEME":   RGBColor(0x09, 0x0D, 0x16),
        "BG_CONTENT": RGBColor(0xFF, 0xFF, 0xFF),
        "TEXT_LIGHT": RGBColor(0xFF, 0xFF, 0xFF),
        "TEXT_DARK":  RGBColor(0x09, 0x0D, 0x16),
        "TEXT_MUTED": RGBColor(0x33, 0x41, 0x55),
        "ACCENT":     RGBColor(0x00, 0x5A, 0x9C),
        "CONTROL":    RGBColor(0x52, 0x52, 0x5B),
        "TINT":       RGBColor(0xF4, 0xF4, 0xF5),
        "ALERT_RED":  RGBColor(0xB9, 0x1C, 0x1C),
        "POS_TEAL":   RGBColor(0x04, 0x78, 0x57),
        "TABLE_ALT":  RGBColor(0xF4, 0xF4, 0xF5),
    },
}

HEAD_FONT = "Calibri"
BODY_FONT = "Calibri"


def create_medical_presentation() -> Presentation:
    """Initializes a 16:9 widescreen PowerPoint presentation."""
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def chars_per_line(w_in: float, pt: float) -> int:
    return max(1, int(w_in * 72 / (pt * CHAR_W_EM) * RAGGED))


def lines_needed(text: str, w_in: float, pt: float) -> int:
    return math.ceil(len(text) / chars_per_line(w_in, pt)) if text else 1


def add_medical_slide(
    prs: Presentation,
    is_bookend: bool = False,
    palette_name: str = "Clinical Navy"
):
    """
    Adds a slide using Title-Only layout (layout[5]) so PowerPoint Outline View
    and accessibility screen readers register the slide title, and sets the
    native slide.background.fill (never a draggable background rectangle).
    """
    pal = PALETTES.get(palette_name, PALETTES["Clinical Navy"])
    layout = prs.slide_layouts[5] if len(prs.slide_layouts) > 5 else prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)

    # Set native slide background fill (cannot be accidentally dragged in PowerPoint)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = pal["BG_THEME"] if is_bookend else pal["BG_CONTENT"]
    return slide


def set_slide_canvas(prs: Presentation, slide, is_bookend: bool = False, palette_name: str = "Clinical Navy"):
    """Sets native slide.background.fill on an existing slide."""
    pal = PALETTES.get(palette_name, PALETTES["Clinical Navy"])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = pal["BG_THEME"] if is_bookend else pal["BG_CONTENT"]
    return slide.background


def add_header(
    slide,
    title_text: str,
    kicker_text: str = "CLINICAL STUDY",
    is_dark: bool = False,
    palette_name: str = "Clinical Navy"
):
    """
    Renders the Title Band (Y=0.45", H=1.50") and synchronizes the native
    title placeholder for PowerPoint Outline View accessibility.
    """
    pal = PALETTES.get(palette_name, PALETTES["Clinical Navy"])
    n = lines_needed(title_text, CONTENT_W, 30)
    assert n <= 2, f"Action title wraps to {n} lines (max 2) -- shorten: {title_text!r}"

    # Populate native title placeholder invisibly off-canvas or reuse it for Outline View
    if slide.shapes.title is not None:
        t_shape = slide.shapes.title
        t_shape.text = title_text
        t_shape.left = Inches(MARGIN_X)
        t_shape.top = Inches(TITLE_TOP)
        t_shape.width = Inches(0.01)
        t_shape.height = Inches(0.01)
        for p in t_shape.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(1)
                r.font.color.rgb = pal["BG_THEME"] if is_dark else pal["BG_CONTENT"]

    tb = slide.shapes.add_textbox(Inches(MARGIN_X), Inches(TITLE_TOP), Inches(CONTENT_W), Inches(TITLE_H))
    tb.name = "Band_Title_Header"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_k = tf.paragraphs[0]
    p_k.alignment = PP_ALIGN.LEFT
    p_k.space_after = Pt(4)
    rk = p_k.add_run()
    rk.text = kicker_text.upper()
    rk.font.size = Pt(13)
    rk.font.name = BODY_FONT
    rk.font.bold = True
    rk.font.color.rgb = pal["ACCENT"] if not is_dark else RGBColor(0x93, 0xC5, 0xFD)

    p_t = tf.add_paragraph()
    p_t.alignment = PP_ALIGN.LEFT
    rt = p_t.add_run()
    rt.text = title_text
    rt.font.size = Pt(30)
    rt.font.name = HEAD_FONT
    rt.font.bold = True
    rt.font.color.rgb = pal["TEXT_LIGHT"] if is_dark else pal["TEXT_DARK"]
    return tb


def add_footer(
    slide,
    citation_text: str,
    slide_num: int,
    is_dark: bool = False,
    lang: str = "id",
    palette_name: str = "Clinical Navy"
):
    """Renders the Footer Band (Y=6.95", H=0.35") with localized slide numbering."""
    pal = PALETTES.get(palette_name, PALETTES["Clinical Navy"])
    tb = slide.shapes.add_textbox(Inches(MARGIN_X), Inches(FOOTER_TOP), Inches(CONTENT_W), Inches(FOOTER_H))
    tb.name = "Band_Footer_Citation"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    page_label = f"Hal. {slide_num}" if lang.lower() == "id" else f"Slide {slide_num}"
    r = p.add_run()
    r.text = f"{citation_text}   |   {page_label}"
    r.font.size = Pt(11.5)
    r.font.name = BODY_FONT
    r.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8) if is_dark else pal["TEXT_MUTED"]
    return tb


def set_bullet(p, char: str = "•", indent_in: float = 0.24):
    """Injects native OpenXML bullet glyph (<a:buChar>) into a paragraph."""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(indent_in)))
    pPr.set("indent", str(-Inches(indent_in)))
    for xml in (
        '<a:buFont %s typeface="Arial"/>' % nsdecls("a"),
        '<a:buChar %s char="%s"/>' % (nsdecls("a"), char),
    ):
        pPr.insert_element_before(parse_xml(xml), "a:buBlip", "a:tabLst", "a:defRPr", "a:extLst")


def add_fragment_bullet(
    tf,
    keyword: str,
    fragment: str,
    box_w_in: float = COL_W_2,
    body_pt: int = 15,
    lead_pt: int = 17,
    is_dark: bool = False,
    bullet_char: str = "•",
    palette_name: str = "Clinical Navy"
):
    """Adds a formatted **Bold Keyword:** telegraphic fragment bullet with line-wrap assertion."""
    pal = PALETTES.get(palette_name, PALETTES["Clinical Navy"])
    full = f"{keyword}: {fragment}"
    n = lines_needed(full, box_w_in, body_pt)
    assert n <= 3, f"Bullet wraps to {n} lines (max 3) -- tighten text: {full!r}"
    p = tf.add_paragraph() if len(tf.paragraphs[0].text) > 0 else tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(10)
    if bullet_char:
        set_bullet(p, char=bullet_char)

    rb = p.add_run()
    rb.text = keyword + (" " if keyword.endswith(":") else ": ")
    rb.font.size = Pt(lead_pt)
    rb.font.name = BODY_FONT
    rb.font.bold = True
    rb.font.color.rgb = pal["TEXT_LIGHT"] if is_dark else pal["TEXT_DARK"]

    rf = p.add_run()
    rf.text = fragment
    rf.font.size = Pt(body_pt)
    rf.font.name = BODY_FONT
    rf.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1) if is_dark else pal["TEXT_MUTED"]
    return p


def add_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    color: Optional[RGBColor] = None,
    palette_name: str = "Clinical Navy"
):
    """Adds a flat, shadow-free, border-free tinted card inside the Body Band."""
    pal = PALETTES.get(palette_name, PALETTES["Clinical Navy"])
    fill_color = color if color is not None else pal["TINT"]
    assert y >= BODY_TOP - 0.05 and (y + h) <= (BODY_TOP + BODY_H + 0.05), (
        f"Box (y={y}, h={h}) violates Body Band ({BODY_TOP}..{BODY_TOP + BODY_H})!"
    )
    box = slide.shapes.add_shape(
        pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    box.name = "Band_Body_Card"
    box.adjustments[0] = 0.05
    box.fill.solid()
    box.fill.fore_color.rgb = fill_color
    box.line.fill.background()
    box.shadow.inherit = False
    for r in box._element.xpath("./p:style/a:effectRef"):
        r.set("idx", "0")
    return box


def add_2_column_body(
    slide,
    left_title: str,
    left_bullets: List[Tuple[str, str]],
    right_title: str,
    right_bullets: List[Tuple[str, str]],
    boxed: bool = False,
    palette_name: str = "Clinical Navy"
):
    """Builds a 2-column comparison or mechanistic layout inside the Body Band."""
    pal = PALETTES.get(palette_name, PALETTES["Clinical Navy"])
    pad = 0.25 if boxed else 0.0
    cols = [(left_title, left_bullets), (right_title, right_bullets)]
    for idx, (col_title, bullets) in enumerate(cols):
        x = MARGIN_X + idx * (COL_W_2 + COL_GAP)
        if boxed:
            add_box(slide, x, BODY_TOP, COL_W_2, BODY_H, pal["TINT"], palette_name)
        tb = slide.shapes.add_textbox(
            Inches(x + pad), Inches(BODY_TOP + pad), Inches(COL_W_2 - 2 * pad), Inches(BODY_H - 2 * pad)
        )
        tb.name = f"Band_Body_Col2_{idx+1}"
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        ph = tf.paragraphs[0]
        ph.space_after = Pt(10)
        rh = ph.add_run()
        rh.text = col_title
        rh.font.size = Pt(20)
        rh.font.name = HEAD_FONT
        rh.font.bold = True
        rh.font.color.rgb = pal["ACCENT"]

        for kw, frag in bullets:
            add_fragment_bullet(
                tf, kw, frag, box_w_in=(COL_W_2 - 2 * pad), body_pt=15, lead_pt=17, palette_name=palette_name
            )


def add_3_column_pillars(
    slide,
    pillars: List[Dict],
    boxed: bool = True,
    palette_name: str = "Clinical Navy"
):
    """
    Builds a 3-pillar layout (e.g. PICO or Trial Schema) inside the Body Band.
    Each item in pillars: {"title": str, "bullets": List[Tuple[str, str]]}
    """
    pal = PALETTES.get(palette_name, PALETTES["Clinical Navy"])
    pad = 0.22 if boxed else 0.0
    assert len(pillars) == 3, "add_3_column_pillars requires exactly 3 pillars."
    for idx, col in enumerate(pillars):
        x = MARGIN_X + idx * (COL_W_3 + COL_GAP)
        if boxed:
            add_box(slide, x, BODY_TOP, COL_W_3, BODY_H, pal["TINT"], palette_name)
        tb = slide.shapes.add_textbox(
            Inches(x + pad), Inches(BODY_TOP + pad), Inches(COL_W_3 - 2 * pad), Inches(BODY_H - 2 * pad)
        )
        tb.name = f"Band_Body_Col3_{idx+1}"
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        ph = tf.paragraphs[0]
        ph.space_after = Pt(10)
        rh = ph.add_run()
        rh.text = col["title"]
        rh.font.size = Pt(20)
        rh.font.name = HEAD_FONT
        rh.font.bold = True
        rh.font.color.rgb = pal["ACCENT"]

        for kw, frag in col["bullets"]:
            add_fragment_bullet(
                tf, kw, frag, box_w_in=(COL_W_3 - 2 * pad), body_pt=14, lead_pt=16, palette_name=palette_name
            )


def add_clinical_table(
    slide,
    headers: List[str],
    rows: List[List[str]],
    x: float = MARGIN_X,
    y: float = BODY_TOP,
    w: float = CONTENT_W,
    h: float = 4.20,
    col_widths: Optional[List[float]] = None,
    palette_name: str = "Clinical Navy"
):
    """
    Creates a styled, publication-grade native PowerPoint table inside the Body Band.
    Supports colorblind-safe status markers ([↑], [↓], [⚠], [✓]).
    """
    pal = PALETTES.get(palette_name, PALETTES["Clinical Navy"])
    assert y >= BODY_TOP - 0.05 and (y + h) <= (BODY_TOP + BODY_H + 0.05), "Table violates Body Band!"
    n_rows = len(rows) + 1
    n_cols = len(headers)
    assert n_cols <= 6 and len(rows) <= 8, f"Table too dense ({n_cols}x{len(rows)}); max 6 cols x 8 rows."

    t_shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(h))
    t_shape.name = "Band_Body_Table"
    table = t_shape.table

    if col_widths and len(col_widths) == n_cols:
        for c_idx, cw in enumerate(col_widths):
            table.columns[c_idx].width = Inches(cw)

    # Style Header Row
    for c_idx, h_text in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = pal["BG_THEME"]
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = h_text
        r.font.name = HEAD_FONT
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = pal["TEXT_LIGHT"]

    # Style Data Rows
    for r_idx, row_data in enumerate(rows, start=1):
        is_alt = (r_idx % 2 == 0)
        for c_idx, val in enumerate(row_data):
            cell = table.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = pal["TABLE_ALT"] if is_alt else pal["BG_CONTENT"]
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = str(val)
            r.font.name = BODY_FONT
            r.font.size = Pt(12.5)
            r.font.bold = (c_idx == 0)
            if "⚠" in str(val) or "[ALERT]" in str(val):
                r.font.color.rgb = pal["ALERT_RED"]
                r.font.bold = True
            elif "✓" in str(val) or "↑" in str(val):
                r.font.color.rgb = pal["POS_TEAL"] if c_idx > 0 else pal["TEXT_DARK"]
            else:
                r.font.color.rgb = pal["TEXT_DARK"]
    return t_shape


def add_native_bar_chart(
    slide,
    categories: List[str],
    series_list: List[Tuple[str, List[float]]],
    x: float,
    y: float,
    w: float,
    h: float,
    palette_name: str = "Clinical Navy"
):
    """Adds an editable native PowerPoint clustered column/bar chart inside the Body Band."""
    pal = PALETTES.get(palette_name, PALETTES["Clinical Navy"])
    assert y >= BODY_TOP - 0.05 and (y + h) <= (BODY_TOP + BODY_H + 0.05), "Chart violates Body Band!"
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for s_name, s_vals in series_list:
        chart_data.add_series(s_name, tuple(s_vals))

    c_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), chart_data
    )
    c_shape.name = "Band_Body_Chart"
    chart = c_shape.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(12)
    chart.legend.font.name = BODY_FONT

    colors = [pal["ACCENT"], pal["CONTROL"], pal["POS_TEAL"], pal["ALERT_RED"]]
    for idx, series in enumerate(chart.series):
        fill = series.format.fill
        fill.solid()
        fill.fore_color.rgb = colors[idx % len(colors)]
    return c_shape


def crop_pdf_figure_300dpi(
    pdf_path: str,
    page_index: int,
    crop_box_norm: Tuple[float, float, float, float],
    out_png: str,
    dpi: int = 300
) -> str:
    """
    Renders a specific page of a medical PDF at 300 DPI using pypdfium2 and crops
    the requested normalized bounding box (left_frac, top_frac, right_frac, bottom_frac).
    """
    doc = pdfium.PdfDocument(pdf_path)
    page = doc[page_index]
    scale = dpi / 72.0
    pil_img = page.render(scale=scale).to_pil()
    w_px, h_px = pil_img.size
    l, t, r, b = crop_box_norm
    cropped = pil_img.crop((int(l * w_px), int(t * h_px), int(r * w_px), int(b * h_px)))
    os.makedirs(os.path.dirname(os.path.abspath(out_png)), exist_ok=True)
    cropped.save(out_png, "PNG")
    doc.close()
    return out_png


def add_cropped_figure(
    slide,
    img_path: str,
    x: float,
    y: float,
    max_w: float,
    max_h: float
):
    """
    Embeds an authentic cropped PDF figure inside the Body Band while preserving
    its native aspect ratio and centering it within (x, y, max_w, max_h).
    """
    assert y >= BODY_TOP - 0.05 and (y + max_h) <= (BODY_TOP + BODY_H + 0.05), "Figure violates Body Band!"
    with Image.open(img_path) as img:
        w_px, h_px = img.size
    aspect = w_px / float(h_px)
    if (max_w / max_h) > aspect:
        final_h = max_h
        final_w = max_h * aspect
    else:
        final_w = max_w
        final_h = max_w / aspect

    offset_x = x + (max_w - final_w) / 2.0
    offset_y = y + (max_h - final_h) / 2.0
    pic = slide.shapes.add_picture(img_path, Inches(offset_x), Inches(offset_y), Inches(final_w), Inches(final_h))
    pic.name = "Band_Body_Figure"
    pic.line.fill.background()
    return pic


def add_speaker_notes(slide, notes_text: str):
    """Adds verbatim clinical traceability notes to the slide's speaker notes pane."""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text
